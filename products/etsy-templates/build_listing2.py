#!/usr/bin/env python3
"""Listing 2: GST/HST Quick Method vs Regular Method Calculator + $30,000 Small-Supplier Tracker (standalone).
Usage: .venv/bin/python products/etsy-templates/build_listing2.py [--out dist]"""
import argparse, pathlib, sys
sys.path.insert(0, str(pathlib.Path(__file__).parent))
from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from build_templates import (PROVINCES, QM_SERVICES, QM_GOODS, QM_GOODS_CREDIT, YESNO, BIZTYPE, H1, H2, B, INPUT_FILL, CALC_FILL, NOTE, BOX, MONEY, PCT, TAX_YEAR, VERSION, header, dv)

def build():
    wb = Workbook()
    L = wb.active; L.title = "Lists"
    header(L, 1, ["Province", "GST", "HST", "PST/QST", "Quick-method PE column"], [26, 8, 8, 10, 22])
    for i, (n, g, h, p, col) in enumerate(PROVINCES, 2):
        L.cell(i, 1, n); L.cell(i, 2, g).number_format = PCT; L.cell(i, 3, h).number_format = PCT; L.cell(i, 4, p).number_format = PCT; L.cell(i, 5, col)
    L["G1"] = "Quick method — services (rows: rate charged 5/13/14/15; cols: PE 5/13/14/15)"; L["G1"].font = B
    for r, row in enumerate(QM_SERVICES, 2):
        L.cell(r, 7, [5, 13, 14, 15][r - 2])
        for c, v in enumerate(row, 8): L.cell(r, c, v).number_format = PCT
    L["G7"] = "Quick method — goods for resale"; L["G7"].font = B
    for r, row in enumerate(QM_GOODS, 8):
        L.cell(r, 7, [5, 13, 14, 15][r - 8])
        for c, v in enumerate(row, 8): L.cell(r, c, v).number_format = PCT
    L["G13"] = "Goods credit on 5% supplies from HST PE"; L["G13"].font = B
    for c, v in enumerate(QM_GOODS_CREDIT, 8): L.cell(14, c, v).number_format = PCT
    L["O1"] = "YesNo"; L["O1"].font = B
    for i, c in enumerate(YESNO, 2): L.cell(i, 15, c)
    L["P1"] = "BizType"; L["P1"].font = B
    for i, c in enumerate(BIZTYPE, 2): L.cell(i, 16, c)
    L["A17"] = f"Source: CRA RC4058 and provincial rates as of 2026-09. Version {VERSION}, tax year {TAX_YEAR}."; L["A17"].font = NOTE

    S = wb.create_sheet("Calculator", 0)
    S.column_dimensions["A"].width = 3; S.column_dimensions["B"].width = 56
    for col in "CDEFG": S.column_dimensions[col].width = 17
    S["B2"] = f"GST/HST Quick Method vs Regular Method — {TAX_YEAR}"; S["B2"].font = H1
    S["B3"] = "Enter your quarterly figures (yellow). The sheet shows what you would remit under each method and whether you must register. Bookkeeping tool, not tax advice."; S["B3"].font = NOTE
    S["B5"] = "1. Settings"; S["B5"].font = H2
    S["B6"] = "Province of your business"; S["C6"] = "Ontario"; dv(S, "Lists!$A$2:$A$14", "C6")
    S["B7"] = "Business type"; S["C7"] = "Services"; dv(S, "Lists!$P$2:$P$3", "C7")
    S["B8"] = "Registered for GST/HST?"; S["C8"] = "Yes"; dv(S, "Lists!$O$2:$O$3", "C8")
    for r in (6, 7, 8): S.cell(r, 3).fill = INPUT_FILL; S.cell(r, 3).border = BOX
    S["B9"] = "GST/HST rate you charge (from province)"; S["C9"] = "=IFERROR(INDEX(Lists!$B$2:$B$14,MATCH($C$6,Lists!$A$2:$A$14,0))+INDEX(Lists!$C$2:$C$14,MATCH($C$6,Lists!$A$2:$A$14,0)),0)"; S["C9"].number_format = PCT
    S["B10"] = "Quick-method PE column"; S["C10"] = "=IFERROR(INDEX(Lists!$E$2:$E$14,MATCH($C$6,Lists!$A$2:$A$14,0)),1)"
    S["B11"] = "Row in the RC4058 table"; S["C11"] = "=IFERROR(MATCH(ROUND(C9*100,0),Lists!$G$2:$G$5,0),1)"
    S["B12"] = "Your RC4058 remittance rate"; S["C12"] = "=IF($C$7=\"Goods for resale\",INDEX(Lists!$H$8:$K$11,C11,C10),INDEX(Lists!$H$2:$K$5,C11,C10))"; S["C12"].number_format = PCT
    S["B13"] = "Goods credit on 5% supplies from an HST province"; S["C13"] = "=IF(AND($C$7=\"Goods for resale\",C11=1),INDEX(Lists!$H$14:$K$14,1,C10),0)"; S["C13"].number_format = PCT
    for r in range(9, 14): S.cell(r, 3).fill = CALC_FILL
    S["B15"] = "2. Your figures by calendar quarter (before tax)"; S["B15"].font = H2
    header(S, 16, ["", "Q1 (Jan–Mar)", "Q2 (Apr–Jun)", "Q3 (Jul–Sep)", "Q4 (Oct–Dec)", "Year"])
    S["B16"] = ""
    S["B17"] = "Taxable sales before GST/HST (exclude zero-rated exports)"
    S["B18"] = "GST/HST you paid on business purchases (ITC-eligible, regular method)"
    S["B19"] = "GST/HST paid on capital purchases (claimable under BOTH methods)"
    for r in (17, 18, 19):
        for c in range(3, 7): S.cell(r, c, 0).number_format = MONEY; S.cell(r, c).fill = INPUT_FILL; S.cell(r, c).border = BOX
        S.cell(r, 7, f"=SUM(C{r}:F{r})").number_format = MONEY; S.cell(r, 7).fill = CALC_FILL
    S["C17"] = 12000; S["D17"] = 15000; S["E17"] = 9000; S["F17"] = 14000; S["C18"] = 260; S["D18"] = 300; S["E18"] = 180; S["F18"] = 320
    S["B20"] = "GST/HST collected (sales × your rate)"
    for c in range(3, 8): S.cell(20, c, f"=ROUND({get_column_letter(c)}17*$C$9,2)").number_format = MONEY; S.cell(20, c).fill = CALC_FILL
    S["B22"] = "3. Regular method"; S["B22"].font = H2
    S["B23"] = "Net tax = collected − ITCs on purchases − ITCs on capital"
    for c in range(3, 8): S.cell(23, c, f"={get_column_letter(c)}20-{get_column_letter(c)}18-{get_column_letter(c)}19").number_format = MONEY; S.cell(23, c).fill = CALC_FILL
    S["B25"] = "4. Quick method"; S["B25"].font = H2
    S["B26"] = "Eligible supplies including GST/HST"
    for c in range(3, 8): S.cell(26, c, f"={get_column_letter(c)}17+{get_column_letter(c)}20").number_format = MONEY; S.cell(26, c).fill = CALC_FILL
    S["B27"] = "Remittance before credit (supplies × RC4058 rate)"
    for c in range(3, 8): S.cell(27, c, f"=ROUND({get_column_letter(c)}26*$C$12,2)").number_format = MONEY; S.cell(27, c).fill = CALC_FILL
    S["B28"] = "1% credit on the first $30,000 of eligible supplies (per fiscal year)"
    S["C28"] = "=ROUND(MIN(C26,30000)*0.01,2)"; S["D28"] = "=ROUND(MAX(0,MIN(C26+D26,30000)-MIN(C26,30000))*0.01,2)"; S["E28"] = "=ROUND(MAX(0,MIN(C26+D26+E26,30000)-MIN(C26+D26,30000))*0.01,2)"; S["F28"] = "=ROUND(MAX(0,MIN(C26+D26+E26+F26,30000)-MIN(C26+D26+E26,30000))*0.01,2)"; S["G28"] = "=SUM(C28:F28)"
    S["B29"] = "Goods credit (goods for resale, 5% supplies from an HST province)"
    for c in range(3, 8): S.cell(29, c, f"=ROUND({get_column_letter(c)}26*$C$13,2)").number_format = MONEY; S.cell(29, c).fill = CALC_FILL
    S["B30"] = "Less ITCs on capital purchases (still claimable)"
    for c in range(3, 8): S.cell(30, c, f"={get_column_letter(c)}19").number_format = MONEY; S.cell(30, c).fill = CALC_FILL
    S["B31"] = "Net tax to remit — quick method"; S["B31"].font = B
    for c in range(3, 8): S.cell(31, c, f"=MAX(0,{get_column_letter(c)}27-{get_column_letter(c)}28-{get_column_letter(c)}29-{get_column_letter(c)}30)").number_format = MONEY; S.cell(31, c).fill = CALC_FILL
    for c in range(3, 8): S.cell(28, c).number_format = MONEY; S.cell(28, c).fill = CALC_FILL
    S["B33"] = "5. Comparison (year)"; S["B33"].font = H2
    S["B34"] = "Regular method — net tax"; S["C34"] = "=G23"
    S["B35"] = "Quick method — net tax"; S["C35"] = "=G31"
    S["B36"] = "You keep more with"; S["C36"] = "=IF(G17=0,\"n/a\",IF(C35<C34,\"Quick method\",\"Regular method\"))"; S["B36"].font = B
    S["B37"] = "Difference per year"; S["C37"] = "=ABS(C34-C35)"
    S["B38"] = "Eligible for the quick method? (worldwide taxable supplies incl. tax ≤ $400,000)"; S["C38"] = "=IF(G26<=400000,\"Likely eligible\",\"Not eligible (over $400,000)\")"
    S["D38"] = "Also excluded: accountants, bookkeepers, lawyers, financial consultants, listed financial institutions, charities and some others (see RC4058). Elect with Form GST74 or My Business Account; stay at least one year."; S["D38"].font = NOTE; S["D38"].alignment = Alignment(wrap_text=True); S.merge_cells("D38:G41")
    for r in (34, 35, 37): S.cell(r, 3).number_format = MONEY
    for r in (34, 35, 36, 37, 38): S.cell(r, 3).fill = CALC_FILL
    S["B43"] = "6. Small-supplier tracker ($30,000 of taxable supplies, before tax)"; S["B43"].font = H2
    S["B44"] = "Highest single calendar quarter (this year)"; S["C44"] = "=MAX(C17:F17)"
    S["B45"] = "Previous year's four quarters, total (enter if you were in business)"; S["C45"] = 0; S["C45"].fill = INPUT_FILL; S["C45"].border = BOX
    S["B46"] = "Rolling four quarters (this year to date; add last year's quarters manually if you started mid-year)"; S["C46"] = "=G17"
    S["B47"] = "Status"; S["C47"] = "=IF($C$8=\"Yes\",\"Registered\",IF(OR(C44>30000,C46>30000,C45>30000),\"Over $30,000: you must register for GST/HST\",IF(C46>25000,\"Approaching $30,000: plan to register\",\"Small supplier (registration optional)\")))"; S["B47"].font = B
    S["D47"] = "Rule: you stop being a small supplier the moment one calendar quarter exceeds $30,000 (register within 29 days), or at the end of the month after four consecutive quarters exceed $30,000."; S["D47"].font = NOTE; S["D47"].alignment = Alignment(wrap_text=True); S.merge_cells("D47:G50")
    for r in (44, 46): S.cell(r, 3).number_format = MONEY
    for r in (44, 46, 47): S.cell(r, 3).fill = CALC_FILL
    S["C45"].number_format = MONEY
    S.conditional_formatting.add("C47", CellIsRule(operator="equal", formula=['"Over $30,000: you must register for GST/HST"'], fill=PatternFill("solid", fgColor="F8D7DA")))
    S["B52"] = "7. Should I register voluntarily while under $30,000?"; S["B52"].font = H2
    S["B53"] = "ITCs you would recover per year (GST/HST paid on purchases + capital)"; S["C53"] = "=G18+G19"
    S["B54"] = "Hours per year you expect to spend on GST/HST filing"; S["C54"] = 6; S["C54"].fill = INPUT_FILL; S["C54"].border = BOX
    S["B55"] = "Value of your time per hour"; S["C55"] = 40; S["C55"].fill = INPUT_FILL; S["C55"].border = BOX
    S["B56"] = "Rough net benefit of registering (ITCs − time cost)"; S["C56"] = "=C53-C54*C55"; S["B56"].font = B
    S["D56"] = "Registering also lets you charge GST/HST to business customers (who recover it) and may look more established; consumers pay more. Once registered you must stay registered at least one year."; S["D56"].font = NOTE; S["D56"].alignment = Alignment(wrap_text=True); S.merge_cells("D56:G59")
    for r in (53, 55, 56): S.cell(r, 3).number_format = MONEY
    for r in (53, 56): S.cell(r, 3).fill = CALC_FILL
    S["B61"] = "This calculator is a bookkeeping tool, not tax advice. Rates checked against CRA RC4058 in September 2026. Verify with the CRA or an accountant before deciding."; S["B61"].font = NOTE
    S["B62"] = f"Version {VERSION} · MapleSheets"; S["B62"].font = NOTE
    S.freeze_panes = "A5"
    return wb

if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--out", default=str(pathlib.Path(__file__).parent / "dist")); a = ap.parse_args()
    out = pathlib.Path(a.out); out.mkdir(parents=True, exist_ok=True)
    p = out / f"GST-HST-Quick-Method-Calculator-Small-Supplier-Tracker-{TAX_YEAR}.xlsx"; build().save(p); print("wrote", p)
