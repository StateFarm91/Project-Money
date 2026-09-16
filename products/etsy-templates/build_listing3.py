#!/usr/bin/env python3
"""Listing 3: Home-Office and Vehicle Expense Workbook (T2125 Part 7 + Chart A), standalone.
Usage: .venv/bin/python products/etsy-templates/build_listing3.py [--out dist]
Reuses the Home office and Vehicle sheet logic from build_templates.py (Listing 1), but the
7N net-income cap becomes a direct input (there is no ledger here to derive it from)."""
import argparse, pathlib, datetime, sys
sys.path.insert(0, str(pathlib.Path(__file__).parent))
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from build_templates import (YESNO, H1, H2, B, INPUT_FILL, CALC_FILL, NOTE, BOX, MONEY, PCT, DATEF, TAX_YEAR, VERSION, header, dv)

ROWS = 200
LOG_FIRST, LOG_LAST = 5, 4 + ROWS

def build():
    wb = Workbook()
    L = wb.active; L.title = "Lists"
    L["A1"] = "YesNo"; L["A1"].font = B
    for i, c in enumerate(YESNO, 2): L.cell(i, 1, c)
    L["A6"] = f"Home-Office and Vehicle Expense Workbook · version {VERSION}, tax year {TAX_YEAR}."; L["A6"].font = NOTE

    # ---------- Start here ----------
    S = wb.create_sheet("Start here", 0)
    S.column_dimensions["A"].width = 3; S.column_dimensions["B"].width = 46; S.column_dimensions["C"].width = 20; S.column_dimensions["D"].width = 60
    S["B2"] = f"Home-Office and Vehicle Expense Workbook — T2125 edition {TAX_YEAR}"; S["B2"].font = H1
    S["B3"] = "Two schedules from Canada's T2125: business-use-of-home (Part 7) and motor vehicle (Chart A). Works in Excel and Google Sheets."; S["B3"].font = NOTE
    S["B5"] = "1. Your net income (from your own T2125 or ledger)"; S["B5"].font = H2
    S["B6"] = "Net income before adjustments (T2125 line 9369, or your best estimate)"; S["C6"] = 20000; S["C6"].fill = INPUT_FILL; S["C6"].border = BOX; S["C6"].number_format = MONEY
    S["D6"] = "The home-office claim cannot exceed this (it can only be carried forward, never create or increase a loss)."; S["D6"].font = NOTE
    S["B8"] = "2. How to use"; S["B8"].font = H2
    for i, s in enumerate(["Home office sheet: enter your space and yearly home costs; the allowable claim (7P) flows from your net income above.",
                            "Vehicle sheet: keep the logbook (the CRA can ask for it); enter total kilometres and yearly costs.",
                            "Both sheets show the exact CRA line the result goes on (9945 and 9281/9936).",
                            "Google Sheets: File > Import > Upload this file > 'Replace spreadsheet'."], 9):
        S.cell(i, 2, f"{i-8}. {s}")
    S["B14"] = "Important"; S["B14"].font = H2
    S["B15"] = (f"This workbook is a bookkeeping tool, not tax, legal or accounting advice. Rates and CRA line numbers were checked in September 2026 for the "
                f"{TAX_YEAR} tax year; verify against current CRA guidance or with an accountant before filing. Keep receipts for six years.")
    S["B15"].alignment = Alignment(wrap_text=True, vertical="top"); S.merge_cells("B15:D17"); S.row_dimensions[15].height = 48
    S["B19"] = f"Version {VERSION} · Built with care (and with AI assistance) in Canada."; S["B19"].font = NOTE

    # ---------- Home office (adapted from build_templates.py) ----------
    Hm = wb.create_sheet("Home office")
    Hm.column_dimensions["A"].width = 3; Hm.column_dimensions["B"].width = 46; Hm.column_dimensions["C"].width = 18; Hm.column_dimensions["D"].width = 60
    Hm["B2"] = "Business-use-of-home expenses (T2125 Part 7)"; Hm["B2"].font = H1
    Hm["B3"] = "Qualify if your home is your principal place of business (>50% of your work) OR the space is used only for business and regularly to meet clients. Enter yearly totals for the whole home."; Hm["B3"].font = NOTE; Hm["B3"].alignment = Alignment(wrap_text=True); Hm.merge_cells("B3:D3"); Hm.row_dimensions[3].height = 40
    Hm["B5"] = "Space"; Hm["B5"].font = H2
    Hm["B6"] = "Area used for business (sq ft or m²)"; Hm["C6"] = 150
    Hm["B7"] = "Total area of home"; Hm["C7"] = 1200
    Hm["B8"] = "Hours per week the space is used for business (24×7=168 if exclusive)"; Hm["C8"] = 168
    Hm["B9"] = "Business-use percentage"; Hm["C9"] = "=IFERROR(MIN(1,C6/C7)*MIN(1,C8/168),0)"; Hm["C9"].number_format = PCT; Hm["C9"].fill = CALC_FILL
    Hm["D9"] = "Area share × time share (time share = 100% if the room is used only for business)."; Hm["D9"].font = NOTE
    Hm["B11"] = "Yearly costs for the whole home"; Hm["B11"].font = H2
    lines = [("7A", "Heat"), ("7B", "Electricity"), ("7C", "Insurance"), ("7D", "Maintenance"), ("7E", "Mortgage interest"), ("7F", "Property taxes"), ("7G", "Other (rent, internet share, condo fees)")]
    for i, (code, lab) in enumerate(lines, 12):
        Hm.cell(i, 2, f"{code} {lab}"); c = Hm.cell(i, 3, 0); c.number_format = MONEY; c.fill = INPUT_FILL; c.border = BOX
    for r in range(6, 9): Hm.cell(r, 3).fill = INPUT_FILL; Hm.cell(r, 3).border = BOX
    Hm["B19"] = "Subtotal"; Hm["C19"] = "=SUM(C12:C18)"
    Hm["B20"] = "Personal-use part"; Hm["C20"] = "=ROUND(C19*(1-C9),2)"
    Hm["B21"] = "7K CCA on the home (business part; usually leave 0 — claiming CCA can affect the principal-residence exemption)"; Hm["C21"] = 0; Hm["C21"].fill = INPUT_FILL; Hm["C21"].border = BOX
    Hm["B22"] = "7L Carry-forward from previous year"; Hm["C22"] = 0; Hm["C22"].fill = INPUT_FILL; Hm["C22"].border = BOX
    Hm["B23"] = "7M Total available (subtotal − personal part + 7K + 7L)"; Hm["C23"] = "=C19-C20+C21+C22"
    Hm["B24"] = "7N Net income after adjustments (from 'Start here'; cannot claim more than this)"; Hm["C24"] = "=MAX(0,'Start here'!$C$6)"
    Hm["B25"] = "7O Amount to carry forward to next year"; Hm["C25"] = "=MAX(0,C23-C24)"
    Hm["B26"] = "7P Allowable claim (lesser of 7M and 7N) → line 9945"; Hm["C26"] = "=MAX(0,MIN(C23,C24))"; Hm["B26"].font = B
    for r in (19, 20, 23, 24, 25, 26): Hm.cell(r, 3).number_format = MONEY; Hm.cell(r, 3).fill = CALC_FILL
    Hm["C21"].number_format = MONEY; Hm["C22"].number_format = MONEY

    # ---------- Vehicle (adapted from build_templates.py; standalone log range) ----------
    V = wb.create_sheet("Vehicle")
    V["A1"] = "Motor vehicle expenses (T2125 Chart A)"; V["A1"].font = H1
    V["A2"] = "Self-employed people deduct ACTUAL costs × business-use %. A cents-per-km claim is not allowed on T2125. Keep a logbook: the CRA can ask for it."; V["A2"].font = NOTE
    header(V, 4, ["Date", "From", "To", "Purpose", "Kilometres", "Business trip?"], [12, 18, 18, 30, 12, 14])
    for r in range(LOG_FIRST, LOG_LAST + 1):
        V.cell(r, 1).number_format = DATEF
        for col in range(1, 7): V.cell(r, col).fill = INPUT_FILL
    dv(V, "Lists!$A$2:$A$3", f"F{LOG_FIRST}:F{LOG_LAST}")
    V.cell(LOG_FIRST, 1, datetime.date(TAX_YEAR, 1, 20)); V.cell(LOG_FIRST, 2, "Home"); V.cell(LOG_FIRST, 3, "Client office"); V.cell(LOG_FIRST, 4, "Example: client meeting (delete me)"); V.cell(LOG_FIRST, 5, 34); V.cell(LOG_FIRST, 6, "Yes")
    V.column_dimensions["H"].width = 44; V.column_dimensions["I"].width = 16
    V["H4"] = "Kilometres"; V["H4"].font = B
    V["H5"] = "Business km (logbook rows marked Yes)"; V["I5"] = f'=SUMIFS(E{LOG_FIRST}:E{LOG_LAST},F{LOG_FIRST}:F{LOG_LAST},"Yes")'
    V["H6"] = "Total km driven this year (all use) — enter"; V["I6"] = 12000; V["I6"].fill = INPUT_FILL; V["I6"].border = BOX
    V["H7"] = "Business-use %"; V["I7"] = "=IFERROR(MIN(1,I5/I6),0)"; V["I7"].number_format = PCT
    V["H9"] = "Yearly vehicle costs (enter)"; V["H9"].font = B
    costs = ["Fuel and oil", "Insurance", "Licence and registration", "Maintenance and repairs", "Interest on vehicle loan (limits apply)", "Leasing costs (limits apply)", "Other (car washes, roadside assistance)"]
    for i, c in enumerate(costs, 10):
        V.cell(i, 8, c); x = V.cell(i, 9, 0); x.number_format = MONEY; x.fill = INPUT_FILL; x.border = BOX
    V["H17"] = "Total vehicle costs"; V["I17"] = "=SUM(I10:I16)"
    V["H18"] = "Deductible part (costs × business %)"; V["I18"] = "=ROUND(I17*I7,2)"
    V["H19"] = "Business parking (100% deductible) — enter"; V["I19"] = 0; V["I19"].fill = INPUT_FILL; V["I19"].border = BOX
    V["H20"] = "Motor vehicle expenses → line 9281"; V["I20"] = "=I18+I19"; V["H20"].font = B
    V["H21"] = "CCA on the vehicle (enter if you claim it) → line 9936"; V["I21"] = 0; V["I21"].fill = INPUT_FILL; V["I21"].border = BOX
    for r in (5, 17, 18, 19, 20, 21): V.cell(r, 9).number_format = MONEY
    for r in (5, 7, 17, 18, 20): V.cell(r, 9).fill = CALC_FILL
    V.freeze_panes = "A5"

    # ---------- Printable logbook ----------
    P = wb.create_sheet("Printable logbook")
    P["A1"] = "Vehicle logbook — printable"; P["A1"].font = H1
    P["A2"] = "Print this sheet and keep it in the car, or fill in the Vehicle sheet directly."; P["A2"].font = NOTE
    header(P, 4, ["Date", "From", "To", "Purpose", "Kilometres", "Business trip? (Y/N)"], [14, 22, 22, 34, 14, 20])
    for r in range(5, 45): P.row_dimensions[r].height = 20
    P.page_setup.orientation = "landscape"; P.print_area = f"A1:F45"

    wb.move_sheet("Lists", offset=len(wb.sheetnames))
    L.sheet_state = "visible"
    return wb

if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--out", default=str(pathlib.Path(__file__).parent / "dist")); a = ap.parse_args()
    out = pathlib.Path(a.out); out.mkdir(parents=True, exist_ok=True)
    p = out / f"Home-Office-Vehicle-Expense-Workbook-{TAX_YEAR}.xlsx"; build().save(p); print("wrote", p)
